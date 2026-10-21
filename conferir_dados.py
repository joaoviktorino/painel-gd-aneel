"""Gera uma planilha Excel de conferência dos resumos (fica fora do Git).

Recalcula os totais direto do parquet bruto, sem usar o processar_dados.py, e
compara com os CSVs de dados/processados/. A planilha vai para dados/brutos/.

Precisa da pasta ferramentas_locais/ (padrão visual Norven, fora do Git).

Uso:  python conferir_dados.py   (rode antes o processar_dados.py)
"""

import json
import sys
from datetime import date, datetime
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import Workbook

PASTA_PROJETO = Path(__file__).resolve().parent
PASTA_BRUTOS = PASTA_PROJETO / "dados" / "brutos"
PASTA_PROCESSADOS = PASTA_PROJETO / "dados" / "processados"
ARQUIVO_PARQUET = PASTA_BRUTOS / "empreendimento-geracao-distribuida.parquet"
PASTA_NORVEN = PASTA_PROJETO / "ferramentas_locais"
LOGO = PASTA_NORVEN / "logo_norven.png"

SIGLA_AUTOR = "JV"
VERSAO = 1
TITULO = "Conferência da base de GD - ANEEL"
SUBTITULO = "Painel de geração distribuída - dados abertos ANEEL"

FORMATO_INTEIRO = "#,##0"
FORMATO_DECIMAL = "#,##0.00"


# ---------- Cálculos direto do bruto ----------

def calcular(con):
    """Recalcula, a partir do parquet, tudo o que a planilha vai mostrar."""
    r = {}
    r["linhas_lidas"] = con.execute("SELECT count(*) FROM bruta").fetchall()[0][0]

    # Mesmas regras do processamento, escritas de novo de forma independente.
    filtro_descarte = "MdaPotenciaInstaladaKW IS NULL OR MdaPotenciaInstaladaKW <= 0"
    r["descartes"] = con.sql(f"""
        SELECT SigUF AS uf, SigAgente AS distribuidora, CodEmpreendimento AS codigo,
               DscFonteGeracao AS fonte, MdaPotenciaInstaladaKW AS potencia_kw,
               DthAtualizaCadastralEmpreend AS data
        FROM bruta WHERE {filtro_descarte} ORDER BY uf, codigo
    """).df()
    r["data_coringa"] = con.execute(f"""
        SELECT count(*) FROM bruta
        WHERE NOT ({filtro_descarte}) AND DthAtualizaCadastralEmpreend < DATE '2000-01-01'
    """).fetchall()[0][0]
    r["fonte_vazia"] = con.execute(
        "SELECT count(*) FROM bruta WHERE DscFonteGeracao IS NULL"
    ).fetchall()[0][0]

    r["por_estado"] = con.sql(f"""
        SELECT SigUF AS uf, any_value(NomRegiao) AS regiao,
               count(*) AS lidas,
               count(*) FILTER (WHERE {filtro_descarte}) AS descartadas,
               sum(MdaPotenciaInstaladaKW) FILTER (WHERE NOT ({filtro_descarte})) AS potencia_kw
        FROM bruta GROUP BY uf ORDER BY uf
    """).df()
    r["por_fonte"] = con.sql(f"""
        SELECT coalesce(DscFonteGeracao, 'Não informada') AS fonte,
               count(*) AS lidas,
               count(*) FILTER (WHERE {filtro_descarte}) AS descartadas,
               coalesce(sum(MdaPotenciaInstaladaKW) FILTER (WHERE NOT ({filtro_descarte})), 0) AS potencia_kw
        FROM bruta GROUP BY fonte ORDER BY fonte
    """).df()

    # Códigos repetidos: exatamente iguais e também ignorando espaços e maiúsculas.
    r["repetidos"] = con.sql("""
        SELECT CodEmpreendimento AS codigo, count(*) AS ocorrencias
        FROM bruta WHERE CodEmpreendimento IS NOT NULL
        GROUP BY codigo HAVING count(*) > 1 ORDER BY ocorrencias DESC, codigo
    """).df()
    r["repetidos_normalizados"] = con.execute("""
        SELECT count(*) FROM (
            SELECT upper(trim(CodEmpreendimento)) FROM bruta
            WHERE CodEmpreendimento IS NOT NULL GROUP BY 1 HAVING count(*) > 1)
    """).fetchall()[0][0]
    r["codigo_vazio"] = con.execute(
        "SELECT count(*) FROM bruta WHERE CodEmpreendimento IS NULL"
    ).fetchall()[0][0]

    periodo, data_geracao = con.execute(
        "SELECT any_value(AnmPeriodoReferencia), max(DatGeracaoConjuntoDados) FROM bruta"
    ).fetchall()[0]
    r["periodo"], r["data_geracao"] = periodo, data_geracao
    return r


# ---------- Montagem da planilha ----------

def formatar_colunas(ws, primeira, ultima, colunas, formato):
    for coluna in colunas:
        for linha in range(primeira, ultima + 1):
            ws[f"{coluna}{linha}"].number_format = formato


def aba_leia_me(wb, nx):
    ws = wb.active
    ws.title = "Leia-me"
    nx.preparar_aba(ws, TITULO, SUBTITULO, str(LOGO))
    linhas = [
        ("Para que serve", "Confere se os resumos publicados no painel (dados/processados/) batem com a base bruta da ANEEL. Os totais foram recalculados direto do arquivo parquet, sem usar o script de processamento."),
        ("Linhas lidas", "Quantidade de registros no arquivo parquet baixado da ANEEL. Cada registro é uma usina (empreendimento de micro ou minigeração)."),
        ("Descartadas", "Linhas que ficam fora de todos os resumos porque a potência instalada é zero, negativa ou vazia. A lista está na aba Descartes."),
        ("Válidas", "Linhas lidas menos descartadas. É o que entra nos totais por estado e por fonte."),
        ("No resumo", "Valor lido do CSV publicado em dados/processados/."),
        ("Diferença", "Fórmula: valor recalculado menos valor do resumo. Tem de ser zero; qualquer outro valor indica erro no processamento."),
        ("Data coringa", "Usinas com data 01/01/1900, um valor de preenchimento. Elas contam nos totais, mas ficam fora do gráfico mensal."),
        ("Código repetido", "Código de empreendimento (CodEmpreendimento) que aparece em mais de uma linha. Em princípio cada usina tem um código único."),
        ("Potência (kW)", "Potência instalada em quilowatts. 1.000 kW = 1 MW; 1.000.000 kW = 1 GW."),
        ("Fonte dos dados", "ANEEL - Relação de empreendimentos de Mini e Micro Geração Distribuída. Licença ODbL."),
        ("Dados pessoais", "A planilha não traz CPF/CNPJ, nome do titular nem CEP. Ela fica em dados/brutos/, que não vai para o GitHub."),
    ]
    nx.tabela(ws, ["Termo", "Explicação"], linhas, [24, 110], colunas_texto=(0, 1))


def aba_resumo(wb, nx, r, ultima_estado):
    ws = wb.create_sheet("Resumo")
    nx.preparar_aba(ws, TITULO, SUBTITULO, str(LOGO))
    soma_estado = f"SUM('Por estado'!{{col}}6:{{col}}{ultima_estado})"
    linhas = [
        ("Linhas lidas do parquet", r["linhas_lidas"], "Total de registros no arquivo bruto."),
        ("Linhas descartadas: potência zero, negativa ou vazia", len(r["descartes"]), "Ficam fora de todos os resumos. Detalhe na aba Descartes."),
        ("Linhas válidas", "=C6-C7", "Fórmula: lidas menos descartadas."),
        ("Usinas nos resumos (soma do CSV por estado)", "=" + soma_estado.format(col="G"), "Soma da coluna 'No resumo' da aba Por estado."),
        ("Diferença: válidas menos resumo", "=C8-C9", "Tem de ser zero."),
        ("Linhas fora só do gráfico mensal: data coringa 01/01/1900", r["data_coringa"], "Contam nos totais por estado e fonte."),
        ("Linhas sem fonte informada", r["fonte_vazia"], "Entram nos resumos como 'Não informada'."),
        ("Códigos de empreendimento repetidos", len(r["repetidos"]), "Códigos que aparecem em mais de uma linha. Lista na aba Códigos repetidos."),
        ("Linhas com código repetido", int(r["repetidos"]["ocorrencias"].sum()) if len(r["repetidos"]) else 0, "Soma das ocorrências dos códigos repetidos."),
        ("Códigos repetidos ignorando espaços e maiúsculas", r["repetidos_normalizados"], "Pega repetições escondidas, como 'GD.SP.1' e 'gd.sp.1 '."),
        ("Linhas sem código de empreendimento", r["codigo_vazio"], ""),
        ("Potência válida recalculada (kW)", "=" + soma_estado.format(col="I"), "Soma da aba Por estado."),
        ("Potência nos resumos (kW)", "=" + soma_estado.format(col="J"), "Soma do CSV por estado."),
        ("Período de referência da base", r["periodo"], "Informado pela ANEEL no próprio arquivo."),
        ("Data de geração da base pela ANEEL", r["data_geracao"], ""),
        ("Data desta conferência", date.today(), ""),
    ]
    ultima = nx.tabela(ws, ["Indicador", "Valor", "Explicação"], linhas, [58, 20, 70],
                       colunas_texto=(0, 2), formulas=True)
    for linha in range(6, ultima + 1):
        valor = ws[f"C{linha}"].value
        if isinstance(valor, (int, float)) or (isinstance(valor, str) and valor.startswith("=")):
            ws[f"C{linha}"].number_format = FORMATO_DECIMAL if "Potência" in ws[f"B{linha}"].value else FORMATO_INTEIRO


def aba_comparacao(wb, nx, nome, chave, rotulo, bruto, resumo, extra=None):
    """Aba com recalculado x resumo, com diferenças em fórmula."""
    ws = wb.create_sheet(nome)
    nx.preparar_aba(ws, TITULO, SUBTITULO, str(LOGO))
    resumo = resumo.set_index(chave)
    cab = [rotulo] + ([extra[1]] if extra else []) + [
        "Linhas lidas", "Descartadas", "Válidas", "No resumo", "Diferença de usinas",
        "Potência válida (kW)", "Potência no resumo (kW)", "Diferença de potência (kW)"]
    linhas = []
    for i, item in enumerate(bruto.itertuples(index=False)):
        r = 6 + i
        chave_valor = getattr(item, chave)
        no_resumo = resumo.loc[chave_valor] if chave_valor in resumo.index else None
        # Letras das colunas: A=#, B=rótulo, (C=extra), depois os números.
        desloc = 1 if extra else 0
        col = {k: chr(ord("C") + desloc + k) for k in range(8)}
        linhas.append(
            [chave_valor] + ([getattr(item, extra[0])] if extra else []) + [
                int(item.lidas), int(item.descartadas),
                f"={col[0]}{r}-{col[1]}{r}",
                int(no_resumo["quantidade"]) if no_resumo is not None else 0,
                f"={col[2]}{r}-{col[3]}{r}",
                round(float(item.potencia_kw), 2),
                float(no_resumo["potencia_kw"]) if no_resumo is not None else 0,
                f"=ROUND({col[5]}{r}-{col[6]}{r},2)",
            ])
    larguras = [22] + ([16] if extra else []) + [14, 14, 14, 14, 14, 18, 18, 18]
    ultima = nx.tabela(ws, cab, linhas, larguras, colunas_texto=(0,), formulas=True)
    desloc = 1 if extra else 0
    inteiros = [chr(ord("C") + desloc + k) for k in range(5)]
    decimais = [chr(ord("C") + desloc + k) for k in range(5, 8)]
    formatar_colunas(ws, 6, ultima, inteiros, FORMATO_INTEIRO)
    formatar_colunas(ws, 6, ultima, decimais, FORMATO_DECIMAL)
    nx.observacao(ws, ultima + 2, "As colunas de diferença são fórmulas e devem ser todas zero.")
    return ultima


def aba_descartes(wb, nx, descartes):
    ws = wb.create_sheet("Descartes")
    nx.preparar_aba(ws, TITULO, SUBTITULO, str(LOGO))
    linhas = [
        [d.uf, d.distribuidora, d.codigo, d.fonte or "Não informada", d.potencia_kw,
         d.data.date() if hasattr(d.data, "date") else d.data,
         "Potência instalada zero, negativa ou vazia"]
        for d in descartes.itertuples(index=False)
    ]
    cab = ["UF", "Distribuidora", "Código do empreendimento", "Fonte", "Potência (kW)",
           "Data de atualização cadastral", "Motivo do descarte"]
    ultima = nx.tabela(ws, cab, linhas, [8, 16, 24, 18, 14, 16, 40], colunas_texto=(6,))
    nx.observacao(ws, ultima + 2, "Sem CPF/CNPJ e sem nome do titular, de propósito.")


def aba_repetidos(wb, nx, repetidos):
    ws = wb.create_sheet("Códigos repetidos")
    nx.preparar_aba(ws, TITULO, SUBTITULO, str(LOGO))
    linhas = [[x.codigo, int(x.ocorrencias)] for x in repetidos.itertuples(index=False)]
    ultima = nx.tabela(ws, ["Código do empreendimento", "Ocorrências"], linhas, [30, 16])
    if not linhas:
        nx.observacao(ws, 7, "Nenhum código de empreendimento aparece em mais de uma linha.")
    else:
        nx.observacao(ws, ultima + 2, "Cada código deveria aparecer uma única vez.")


def main():
    for arquivo in (ARQUIVO_PARQUET, PASTA_PROCESSADOS / "por_estado.csv", LOGO):
        if not arquivo.exists():
            raise FileNotFoundError(f"Falta {arquivo.relative_to(PASTA_PROJETO)}.")
    sys.path.insert(0, str(PASTA_NORVEN))
    import norven_xlsx as nx

    con = duckdb.connect()
    con.execute(f"CREATE VIEW bruta AS SELECT * FROM read_parquet('{ARQUIVO_PARQUET.as_posix()}')")
    r = calcular(con)
    resumo_estado = pd.read_csv(PASTA_PROCESSADOS / "por_estado.csv")
    resumo_fonte = pd.read_csv(PASTA_PROCESSADOS / "por_fonte.csv")

    wb = Workbook()
    aba_leia_me(wb, nx)
    # "Por estado" é criada antes do Resumo para sabermos a última linha, e
    # depois reordenada para o Resumo ficar logo após o Leia-me.
    ultima_estado = aba_comparacao(wb, nx, "Por estado", "uf", "UF", r["por_estado"],
                                   resumo_estado, extra=("regiao", "Região"))
    aba_comparacao(wb, nx, "Por fonte", "fonte", "Fonte", r["por_fonte"], resumo_fonte)
    aba_descartes(wb, nx, r["descartes"])
    aba_repetidos(wb, nx, r["repetidos"])
    aba_resumo(wb, nx, r, ultima_estado)
    wb.move_sheet("Resumo", offset=-(len(wb.sheetnames) - 2))
    nx.formatar_datas(wb)

    nome = nx.nome_arquivo(datetime.now().strftime("%y%m%d"), VERSAO, SIGLA_AUTOR,
                           "Conferência da base de GD")
    destino = PASTA_BRUTOS / nome
    wb.save(destino)
    print(f"Planilha salva em {destino.relative_to(PASTA_PROJETO)}")


if __name__ == "__main__":
    try:
        main()
    except Exception as erro:
        print(f"Erro: {erro}", file=sys.stderr)
        sys.exit(1)
