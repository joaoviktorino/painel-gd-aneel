"""Investiga se DthAtualizaCadastralEmpreend serve como data de conexão (issue #1).

Lê o parquet bruto e grava em dados/brutos/ (fora do Git):
  - investigacao_datas/por_mes.csv        quantidade e potência por mês, desde 2012
  - investigacao_datas/top_dias.csv       os 10 dias com mais registros
  - investigacao_datas/mes_porte.csv      mês x porte (micro/minigeração)
  - investigacao_datas/mes_fonte.csv      mês x fonte de geração
  - uma planilha Excel com esses resultados, no padrão visual Norven

A planilha precisa da pasta ferramentas_locais/ (fora do Git). Sem ela, o
script gera só os CSVs.

Uso:  python investigar_datas.py   (rode antes o baixar_dados.py)
"""

import sys
from datetime import date, datetime

from checar_dados import PASTA_BRUTOS, PASTA_PROJETO, abrir_bruta

PASTA_SAIDA = PASTA_BRUTOS / "investigacao_datas"
PASTA_NORVEN = PASTA_PROJETO / "ferramentas_locais"
LOGO = PASTA_NORVEN / "logo_norven.png"

DATA_INICIAL = "2012-01-01"
QUANTIDADE_DIAS = 10

SIGLA_AUTOR = "JV"
VERSAO = 1
TITULO = "Investigação da data de conexão - GD ANEEL"
SUBTITULO = "Painel de geração distribuída - issue #1 (DthAtualizaCadastralEmpreend)"

# Alguns valores da ANEEL vêm sem acento.
ROTULOS = {"Microgeracao": "Microgeração", "Minigeracao": "Minigeração"}

# Mesmas regras do painel: só usinas com potência positiva e datas a partir de
# DATA_INICIAL (o que exclui também as datas coringa de 1900).
SQL_BASE = f"""
    CREATE TABLE base AS
    SELECT
        DthAtualizaCadastralEmpreend                         AS data,
        CAST(date_trunc('month', DthAtualizaCadastralEmpreend) AS DATE) AS mes,
        trim(DscPorte)                                       AS porte,
        coalesce(DscFonteGeracao, 'Não informada')           AS fonte,
        MdaPotenciaInstaladaKW                               AS potencia_kw
    FROM bruta
    WHERE MdaPotenciaInstaladaKW > 0
      AND DthAtualizaCadastralEmpreend >= DATE '{DATA_INICIAL}'
"""

SQL_POR_MES = """
    SELECT mes, count(*) AS quantidade, round(sum(potencia_kw), 2) AS potencia_kw
    FROM base GROUP BY mes ORDER BY mes
"""

SQL_TOP_DIAS = f"""
    SELECT data, count(*) AS quantidade, round(sum(potencia_kw), 2) AS potencia_kw
    FROM base GROUP BY data ORDER BY quantidade DESC, data LIMIT {QUANTIDADE_DIAS}
"""

SQL_MES_POR = """
    SELECT mes, {coluna}, count(*) AS quantidade, round(sum(potencia_kw), 2) AS potencia_kw
    FROM base GROUP BY ALL ORDER BY mes, {coluna}
"""


# ---------- Cálculos ----------

def calcular():
    con = abrir_bruta()
    con.execute(SQL_BASE)
    resultados = {
        "por_mes": con.sql(SQL_POR_MES).df(),
        "top_dias": con.sql(SQL_TOP_DIAS).df(),
        "mes_porte": con.sql(SQL_MES_POR.format(coluna="porte")).df(),
        "mes_fonte": con.sql(SQL_MES_POR.format(coluna="fonte")).df(),
    }
    resultados["mes_porte"]["porte"] = resultados["mes_porte"]["porte"].replace(ROTULOS)
    # Datas como data simples (sem hora), para o CSV e o Excel.
    for tabela in resultados.values():
        for coluna in ("mes", "data"):
            if coluna in tabela:
                tabela[coluna] = tabela[coluna].dt.date
    return resultados


def salvar_csvs(resultados):
    PASTA_SAIDA.mkdir(parents=True, exist_ok=True)
    for nome, tabela in resultados.items():
        tabela.to_csv(PASTA_SAIDA / f"{nome}.csv", index=False)
        print(f"{PASTA_SAIDA.relative_to(PASTA_PROJETO) / (nome + '.csv')}: {len(tabela)} linhas")


# ---------- Planilha ----------

def formatar_numeros(ws, primeira, ultima, colunas, formato):
    for coluna in colunas:
        for linha in range(primeira, ultima + 1):
            ws[f"{coluna}{linha}"].number_format = formato


def aba_leia_me(wb, nx):
    ws = wb.active
    ws.title = "Leia-me"
    nx.preparar_aba(ws, TITULO, SUBTITULO, str(LOGO))
    linhas = [
        ("Para que serve", "Ajuda a verificar se a data DthAtualizaCadastralEmpreend (data da última atualização cadastral) pode ser usada como data de conexão das usinas, como o painel faz hoje. Ver issue #1 no GitHub."),
        ("Base usada", f"Arquivo parquet da ANEEL. Só usinas com potência instalada maior que zero e com data a partir de {date.fromisoformat(DATA_INICIAL):%d/%m/%Y} (o que exclui as datas de preenchimento 01/01/1900)."),
        ("Por mês", "Quantidade de usinas e potência somada em cada mês da data de atualização cadastral. A coluna Mês traz o primeiro dia do mês."),
        ("10 dias com mais registros", "Os dias com mais usinas com essa data. Muitos registros num único dia podem indicar atualização em lote (por exemplo, na migração do SISGD para o MMGD), e não conexões reais."),
        ("Mês x porte", "Usinas e potência de cada mês, separadas por porte (microgeração e minigeração). A coluna Total é uma fórmula."),
        ("Mês x fonte", "Usinas de cada mês, por fonte de geração. Fontes que eram mais comuns em usinas antigas aparecendo com data recente podem indicar usinas antigas com cadastro atualizado. A coluna Total é uma fórmula."),
        ("Suspensão da base", "A ANEEL suspendeu a atualização da base entre 23/09/2025 e 13/11/2025, durante a migração do SISGD para o MMGD."),
        ("Potência (kW)", "Potência instalada em quilowatts. 1.000 kW = 1 MW."),
        ("Dados pessoais", "A planilha não traz CPF/CNPJ, nome do titular nem CEP. Ela fica em dados/brutos/, fora do GitHub."),
        ("Fonte dos dados", "ANEEL - Relação de empreendimentos de Mini e Micro Geração Distribuída. Licença ODbL."),
    ]
    nx.tabela(ws, ["Termo", "Explicação"], linhas, [26, 110], colunas_texto=(0, 1))


def aba_simples(wb, nx, nome, tabela, coluna_data, rotulo_data):
    ws = wb.create_sheet(nome)
    nx.preparar_aba(ws, TITULO, SUBTITULO, str(LOGO))
    linhas = [[getattr(r, coluna_data), int(r.quantidade), float(r.potencia_kw)] for r in tabela.itertuples()]
    ultima = nx.tabela(ws, [rotulo_data, "Usinas", "Potência (kW)"], linhas, [16, 14, 18])
    formatar_numeros(ws, 6, ultima, ["C"], "#,##0")
    formatar_numeros(ws, 6, ultima, ["D"], "#,##0.00")
    return ws, ultima


def aba_cruzada(wb, nx, nome, tabela, coluna, valores, incluir_potencia):
    """Uma linha por mês e uma coluna por categoria, com Total em fórmula."""
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(nome)
    nx.preparar_aba(ws, TITULO, SUBTITULO, str(LOGO))
    usinas = tabela.pivot_table(index="mes", columns=coluna, values="quantidade", aggfunc="sum", fill_value=0)
    usinas = usinas.reindex(columns=valores, fill_value=0)
    cabecalhos = ["Mês (1º dia)"] + [f"{v} - usinas" if incluir_potencia else v for v in valores] + ["Total de usinas"]
    if incluir_potencia:
        potencia = tabela.pivot_table(index="mes", columns=coluna, values="potencia_kw", aggfunc="sum", fill_value=0)
        potencia = potencia.reindex(columns=valores, fill_value=0)
        cabecalhos += [f"{v} - potência (kW)" for v in valores] + ["Total de potência (kW)"]

    # Coluna A = '#', B = mês; as categorias começam na C.
    n = len(valores)
    primeira_usinas, ultima_usinas = get_column_letter(3), get_column_letter(2 + n)
    linhas = []
    for i, mes in enumerate(usinas.index):
        r = 6 + i
        linha = [mes] + [int(x) for x in usinas.loc[mes]] + [f"=SUM({primeira_usinas}{r}:{ultima_usinas}{r})"]
        if incluir_potencia:
            ini, fim = get_column_letter(4 + n), get_column_letter(3 + 2 * n)
            linha += [round(float(x), 2) for x in potencia.loc[mes]] + [f"=SUM({ini}{r}:{fim}{r})"]
        linhas.append(linha)

    rotacionar = () if incluir_potencia else tuple(range(1, n + 1))
    larguras = [16] + [18 if incluir_potencia else 6] * n + [16]
    if incluir_potencia:
        larguras += [20] * n + [20]
    ultima = nx.tabela(ws, cabecalhos, linhas, larguras, formulas=True, colunas_rotacionadas=rotacionar)
    total_colunas = len(cabecalhos) + 1
    inteiros = [get_column_letter(c) for c in range(3, 4 + n)]
    formatar_numeros(ws, 6, ultima, inteiros, "#,##0")
    if incluir_potencia:
        decimais = [get_column_letter(c) for c in range(4 + n, total_colunas + 1)]
        formatar_numeros(ws, 6, ultima, decimais, "#,##0.00")
    return ws, ultima


def gerar_planilha(resultados):
    if not LOGO.exists():
        print("Aviso: ferramentas_locais/ não encontrada; planilha Excel não gerada (só os CSVs).")
        return
    sys.path.insert(0, str(PASTA_NORVEN))
    import norven_xlsx as nx
    from openpyxl import Workbook

    wb = Workbook()
    aba_leia_me(wb, nx)
    aba_simples(wb, nx, "Por mês", resultados["por_mes"], "mes", "Mês (1º dia)")
    aba_simples(wb, nx, "10 dias com mais registros", resultados["top_dias"], "data", "Data")

    portes = resultados["mes_porte"].groupby("porte")["quantidade"].sum().sort_values(ascending=False).index.tolist()
    aba_cruzada(wb, nx, "Mês x porte", resultados["mes_porte"], "porte", portes, incluir_potencia=True)
    fontes = resultados["mes_fonte"].groupby("fonte")["quantidade"].sum().sort_values(ascending=False).index.tolist()
    ws, ultima = aba_cruzada(wb, nx, "Mês x fonte", resultados["mes_fonte"], "fonte", fontes, incluir_potencia=False)
    nx.observacao(ws, ultima + 2, "Cabeçalhos girados para caber todas as fontes. Fontes em ordem decrescente de usinas.")

    nx.formatar_datas(wb)
    nome = nx.nome_arquivo(datetime.now().strftime("%y%m%d"), VERSAO, SIGLA_AUTOR,
                           "Investigação da data de conexão")
    destino = PASTA_BRUTOS / nome
    wb.save(destino)
    print(f"Planilha salva em {destino.relative_to(PASTA_PROJETO)}")


def main():
    resultados = calcular()
    salvar_csvs(resultados)
    gerar_planilha(resultados)


if __name__ == "__main__":
    try:
        main()
    except Exception as erro:
        print(f"Erro: {erro}", file=sys.stderr)
        sys.exit(1)
